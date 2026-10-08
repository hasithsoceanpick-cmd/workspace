import os
_HERE = os.path.dirname(os.path.abspath(__file__))
os.makedirs(os.path.join(_HERE, 'shots'), exist_ok=True)
# Browser tests: deadline history on cards, repeating tasks, "Who's on what" board,
# weekly summary, Excel exports, older logins, Month-end declaration (Finance only).
import asyncio, datetime, subprocess
import openpyxl
from playwright.async_api import async_playwright

BASE = os.environ.get('APP_URL', 'http://localhost:4173')
DB = os.environ.get('DATABASE_URL', 'postgres://postgres@127.0.0.1:54322/postgres')
OUT = os.path.join(_HERE, 'shots')
results, errors = [], []
def check(c, m): results.append(c); print(('PASS ' if c else 'FAIL ') + m, flush=True)
def sql(q): return subprocess.run(['psql', DB, '-At', '-c', q], capture_output=True, text=True).stdout.strip()
def day(n): return (datetime.date.today() + datetime.timedelta(days=n)).isoformat()

async def page_for(b, email, w=1360, h=860, pw='password1'):
    ctx = await b.new_context(viewport={'width': w, 'height': h}, timezone_id='Asia/Colombo', accept_downloads=True)
    pg = await ctx.new_page()
    pg.on('pageerror', lambda e: errors.append(f'{email}: {e}'))
    pg.on('console', lambda m: errors.append(f'{email}: {m.text}') if m.type == 'error' and 'Failed to load resource' not in m.text else None)
    await pg.goto(BASE); await pg.fill('input[type=email]', email); await pg.fill('input[type=password]', pw)
    await pg.click('button:has-text("Sign in")'); await pg.wait_for_timeout(1500)
    return pg

async def to_finance(h):
    if 'Finance' not in await h.locator('.dept-switch .switcher').inner_text():
        await h.click('.dept-switch .switcher'); await h.click('.drop-item:has-text("Finance")'); await h.wait_for_timeout(1000)

async def saved(dl_info, name):
    path = os.path.join(OUT, name)
    await (await dl_info.value).save_as(path)
    return path

async def toast(pg):
    try:
        await pg.wait_for_selector('.toast', timeout=4000)
        return ' | '.join(await pg.locator('.toast').all_inner_texts())
    except Exception:
        return ''

async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()
        h = await page_for(b, 'hasith@demo.lk')
        await to_finance(h)

        # ---------- deadline history on the card ----------
        await h.goto(BASE + '/#/tasks/list?who=all'); await h.wait_for_timeout(1000)
        row = h.locator('.task-row:has-text("WHT schedule for September")')
        check(await row.locator('.pill.moved').inner_text() == 'Moved 1×', 'a moved task shows "Moved 1×" on its card')
        check(await row.locator('.tr-due .was').count() == 1, 'the first deadline is shown struck through')
        await row.locator('.pill.moved').click(); await h.wait_for_timeout(700)
        pop = await h.locator('.due-pop').inner_text()
        check('Amaya moved it' in pop and 'first deadline' in pop, f'clicking it shows who moved it and from/to ({pop[:80]!r})')
        check(await h.locator('.drawer').count() == 0, 'opening the history does not open the task')
        await h.keyboard.press('Escape'); await h.wait_for_timeout(200)
        await row.click(); await h.wait_for_selector('.drawer')
        await h.click('.due-history .link'); await h.wait_for_timeout(300)
        check('moved it' in await h.locator('.due-history').inner_text(), 'the task panel shows the deadline history too')
        await h.screenshot(path=f'{OUT}/r1_due_history.png')
        await h.keyboard.press('Escape')

        # ---------- repeating task ----------
        await h.click('.subnav button:has-text("New task")'); await h.wait_for_selector('.drawer')
        await h.fill('.title-input', 'VAT return'); await h.fill('.drawer input[type=date]', day(3))
        await h.select_option('.drawer .field select >> nth=0', label='Kasun Silva')
        await h.select_option('.drawer select:near(:text("Repeat"))', label='Every month')
        await h.fill('.step-add input', 'Reconcile input VAT'); await h.press('.step-add input', 'Enter')
        await h.click('button:has-text("Create task")'); await h.wait_for_timeout(1500)
        vrow = h.locator('.task-row:has(.tr-title:text-is("VAT return"))')
        check(await vrow.locator('.pill.repeat').count() == 1, 'a repeating task shows ⟳ Monthly')
        await vrow.locator('.check').click(); await h.wait_for_timeout(900)
        t = await toast(h)
        check('next one has been created' in t, f'completing it says the next one was created ({t})')
        await h.wait_for_timeout(1500)
        nxt = sql("select t2.due_date, t2.status, (select count(*) from task_checklist c where c.task_id=t2.id) from tasks t join tasks t2 on t2.id=t.next_task_id where t.title='VAT return'")
        d3 = datetime.date.fromisoformat(day(3)); m = d3.month % 12 + 1; y = d3.year + (1 if d3.month == 12 else 0)
        import calendar
        exp = datetime.date(y, m, min(d3.day, calendar.monthrange(y, m)[1])).isoformat()
        check(nxt.startswith(f'{exp}|todo|1'), f'the next one is due a month later with its checklist ({nxt} vs {exp})')
        await h.goto(BASE + '/#/tasks/list?who=all'); await h.wait_for_timeout(1000)
        check(await h.locator('.task-row:not(.done):has(.tr-title:text-is("VAT return"))').count() == 1, 'and it appears in the list')

        # ---------- Excel export from the list ----------
        async with h.expect_download() as dl:
            await h.click('button:has-text("Export to Excel")')
        wb = openpyxl.load_workbook(await saved(dl, 'tasks.xlsx'))
        check(wb.sheetnames == ['Tasks', 'Deadline moves'], f'export has Tasks and Deadline moves sheets ({wb.sheetnames})')
        ws = wb['Tasks']
        heads = [c.value for c in ws[1]]
        n_tasks = int(sql("select count(*) from tasks t join departments d on d.id=t.department_id where d.name='Finance'"))
        check(ws.max_row - 1 == n_tasks and heads[:3] == ['Task #', 'Title', 'Owner'], f'every Finance task is exported ({ws.max_row - 1} of {n_tasks})')
        wht = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == 'WHT schedule for September'][0]
        check(isinstance(wht[6], datetime.datetime) and wht[8] == 1, 'deadlines are real dates and moves are counted')
        check(ws.freeze_panes == 'A2', 'header row is frozen')
        mv = wb['Deadline moves']
        check(mv.max_row >= 2 and mv.cell(2, 2).value == 'WHT schedule for September', 'the moves sheet lists who moved what')

        # ---------- Who's on what ----------
        await h.goto(BASE + '/#/tasks/team'); await h.wait_for_timeout(1200)
        check(await h.locator('.tab.active').inner_text() == "Who's on what", "the Team tab is now 'Who's on what'")
        cols = await h.locator('.bcol-name').all_inner_texts()
        check(len(cols) == 4 and any('Kasun' in c for c in cols), f'one column per person ({cols})')
        kas = h.locator('.bcol:has(.bcol-name:has-text("Kasun"))')
        ama = h.locator('.bcol:has(.bcol-name:has-text("Amaya"))')
        check(await kas.locator('.bcard:has-text("Petty cash count")').count() == 1, "Kasun's column shows his open work")
        await kas.locator('.bcard:has-text("Fixed asset register update")').drag_to(ama.locator('.bcol-body')); await h.wait_for_timeout(1200)
        t = await toast(h)
        check('Handed to Amaya' in t and await ama.locator('.bcard:has-text("Fixed asset register update")').count() == 1, f'drag a card to hand it over ({t})')
        await h.screenshot(path=f'{OUT}/r2_board.png')
        await h.select_option('.filters select', 'overdue'); await h.wait_for_timeout(500)
        late_cards = await h.locator('.bcard').count()
        check(late_cards == await h.locator('.bcard.late').count() and late_cards > 0, 'overdue filter shows only late work')
        await h.click('.seg-btn:has-text("Table")'); await h.wait_for_timeout(500)
        check('Deadline moved' in await h.locator('.team-table thead').inner_text(), 'table view has a deadline-moved column')

        # ---------- Weekly summary ----------
        await h.goto(BASE + '/#/tasks/week'); await h.wait_for_timeout(1000)
        check('Weekly summary' in await h.locator('h1').inner_text(), 'weekly summary page opens')
        await h.click('button:has-text("This week")'); await h.wait_for_timeout(1000)
        body = await h.locator('.week-lists').inner_text()
        check('WHT schedule for September' in body, 'this week: the moved deadline is listed')
        kas_row = await h.locator('.team-table tr:has-text("Kasun")').inner_text()
        check(kas_row.split('\t')[1].strip() not in ('0', ''), f'Kasun shows work done this week ({kas_row!r})')
        async with h.expect_download() as dl:
            await h.click('button:has-text("Export to Excel")')
        wb = openpyxl.load_workbook(await saved(dl, 'week.xlsx'))
        check(wb.sheetnames == ['Summary', 'Deadlines moved', 'Completed late', 'Overdue now'], f'weekly export sheets ({wb.sheetnames})')
        await h.screenshot(path=f'{OUT}/r3_week.png')
        k = await page_for(b, 'kasun@demo.lk')
        check(await k.locator('.tab:has-text("Weekly summary")').count() == 0 and await k.locator(".tab:has-text(\"Who's on what\")").count() == 0, 'members do not get the board or the summary')

        # ---------- older login without a profile ----------
        uid = sql("insert into auth.users (email, encrypted_password, raw_user_meta_data) values ('zumra@demo.lk','password1','{\"full_name\":\"Zumra Ahamed\"}') returning id").splitlines()[0]
        sql(f"delete from profiles where id='{uid}'; delete from notifications where kind='signup'")
        z = await page_for(b, 'zumra@demo.lk')
        check(await z.locator('h1:has-text("Waiting for approval")').count() == 1, 'an older login lands on "Waiting for approval", not "No access"')
        await h.goto(BASE + '/#/admin/people'); await h.wait_for_timeout(1200)
        check('Zumra Ahamed' in await h.locator('main').inner_text(), 'and shows up in Admin → People to approve')

        # ---------- Month-end declaration ----------
        await h.goto(BASE + '/#/admin/departments'); await h.wait_for_timeout(1000)
        fin = h.locator('.dept-card').filter(has_text='Finance').first
        await fin.locator('.feature-row:has-text("Month-end declaration") input').click(); await h.wait_for_timeout(1000)
        hr = h.locator('.dept-card').filter(has_text='HR').first
        check(not await hr.locator('.feature-row:has-text("Month-end declaration") input').is_checked(), 'it stays off for HR')
        n = await page_for(b, 'nadeesha@demo.lk')
        await n.goto(BASE + '/#/tasks/x-month-end'); await n.wait_for_timeout(1200)
        check(await n.locator('.tab.active').inner_text() == 'Month end', 'Finance gets a Month end tab')
        await n.click('button:has-text("Edit master list")'); await n.wait_for_timeout(600)
        for code, title, owner in [('MEC-01', 'All bank accounts reconciled', 'Kasun Silva'), ('MEC-02', 'Petty cash counted and signed', 'Amaya Fernando'), ('CMP-01', 'VAT and WHT returns filed', 'Nadeesha Perera')]:
            await n.fill('.me-add input[aria-label="New code"]', code)
            await n.fill('.me-add input[aria-label="New category"]', code[:3])
            await n.fill('.me-add input[aria-label="New line"]', title)
            await n.select_option('.me-add select[aria-label="New owner"]', label=owner)
            await n.click('.me-add button:has-text("Add line")'); await n.wait_for_timeout(600)
        check(await n.locator('.me-table tbody tr').count() == 3, 'a senior builds the master list')
        await n.click('button:has-text("Back to the month")'); await n.wait_for_timeout(600)
        last_month = (datetime.date.today().replace(day=1) - datetime.timedelta(days=1)).strftime('%B %Y')
        await n.click(f'button:has-text("Start {last_month}")'); await n.wait_for_timeout(1500)
        check(await n.locator('.me-row').count() == 3, f'starting {last_month} copies the lines')
        check('0/3 lines ticked' in await n.locator('.me-status').inner_text(), 'progress starts at 0/3')
        await k.goto(BASE + '/#/tasks/x-month-end?who=me'); await k.reload(); await k.wait_for_timeout(1500)
        check(await k.locator('.me-row').count() == 1, "'Mine' shows only my lines")
        await k.locator('.me-row input[type=checkbox]').check(); await k.wait_for_timeout(700)
        await k.fill('.me-remarks', 'All 6 accounts agree to statements'); await k.press('.me-remarks', 'Enter'); await k.wait_for_timeout(700)
        a = await page_for(b, 'amaya@demo.lk')
        await a.goto(BASE + '/#/tasks/x-month-end'); await a.wait_for_timeout(1200)
        check(await a.locator('.me-row:has-text("All bank accounts") input[type=checkbox]').is_disabled(), "others can't tick someone else's line")
        check('All 6 accounts' in await a.locator('.me-row:has-text("All bank accounts") .me-remarks').input_value(), 'remarks are saved and visible')
        await n.reload(); await n.wait_for_timeout(1200)
        await n.locator('.me-row:has-text("VAT and WHT") input[type=checkbox]').check(); await n.wait_for_timeout(1000)
        check(await n.locator('button:has-text("Mark reviewed")').is_disabled(), "can't review until every line is ticked")
        await a.locator('.me-row:has-text("Petty cash") input[type=checkbox]').check(); await a.wait_for_timeout(1000)
        await n.reload(); await n.wait_for_timeout(1200)
        await n.click('.bell .icon-btn'); await n.wait_for_timeout(300)
        check('ready for review' in await n.locator('.bell-pop').inner_text(), 'last tick tells the senior it is ready for review')
        await n.keyboard.press('Escape'); await n.mouse.click(10, 500)
        await n.click('button:has-text("Mark reviewed")'); await n.wait_for_timeout(1000)
        check(await n.locator('.me-badge').inner_text() == 'Reviewed', 'the senior reviews')
        check(await n.locator('button:has-text("Approve")').count() == 0, 'only the manager sees Approve')
        await h.reload(); await h.wait_for_timeout(1500)
        await h.click('.bell .icon-btn'); await h.wait_for_timeout(300)
        await h.click('.notice-row:has-text("ready for your approval")'); await h.wait_for_timeout(1800)
        check('x-month-end' in h.url, 'the alert opens the Month end page')
        await h.click('button:has-text("Approve")'); await h.wait_for_timeout(1000)
        check(await h.locator('.me-badge').inner_text() == 'Approved', 'the manager approves')
        await k.reload(); await k.wait_for_timeout(1200)
        check(await k.locator('.me-row input[type=checkbox]').is_disabled(), 'after sign-off the lines are locked')
        await h.screenshot(path=f'{OUT}/r4_month_end.png')
        async with h.expect_download() as dl:
            await h.click('.me-actions button:has-text("Export to Excel")')
        wb = openpyxl.load_workbook(await saved(dl, 'month-end.xlsx'))
        check(wb.sheetnames[1] == 'Sign-off' and wb.worksheets[0].max_row == 4, f'month export has the lines and the sign-off ({wb.sheetnames})')
        pr = await page_for(b, 'priya@demo.lk')
        check(await pr.locator('.tab:has-text("Month end")').count() == 0, 'HR has no Month end tab')
        await pr.goto(BASE + '/#/tasks/x-month-end'); await pr.wait_for_timeout(1000)
        check(await pr.locator('.me-row').count() == 0, 'and cannot open it by link')

        # phone layout
        m = await page_for(b, 'hasith@demo.lk', 390, 820)
        await to_finance(m)
        for path in ['/#/tasks/team', '/#/tasks/week', '/#/tasks/x-month-end']:
            await m.goto(BASE + path); await m.wait_for_timeout(1000)
            sw = await m.evaluate('document.documentElement.scrollWidth')
            check(sw <= 392, f'{path} fits a phone (scrollWidth {sw})')
        await m.screenshot(path=f'{OUT}/r5_month_end_phone.png', full_page=True)
        await b.close()

    print(f'\n{sum(results)} passed, {len(results) - sum(results)} failed')
    errs = [e for e in errors]
    if errs: print('Browser errors:\n  ' + '\n  '.join(errs[:15]))

asyncio.run(main())
