import os
_HERE = os.path.dirname(os.path.abspath(__file__))
os.makedirs(os.path.join(_HERE, 'shots'), exist_ok=True)
# Browser tests for round 5: install as an app (PWA), phone alerts end to end (with fake phones),
# checker sign-off, "Got it", early reminders, Compliance calendar, Trends and Ctrl+K search.
import asyncio, datetime, json, subprocess, urllib.parse, urllib.request
import openpyxl
from playwright.async_api import async_playwright

BASE = os.environ.get('APP_URL', 'http://localhost:4173')
MOCK = os.environ.get('MOCK_URL', 'http://localhost:54321')
DB = os.environ.get('DATABASE_URL', 'postgres://postgres@127.0.0.1:54322/postgres')
OUT = os.path.join(_HERE, 'shots')
results, errors = [], []
def check(c, m): c = bool(c); results.append(c); print(('PASS ' if c else 'FAIL ') + m, flush=True)
def sql(q): return subprocess.run(['psql', DB, '-At', '-c', q], capture_output=True, text=True).stdout.strip()
def day(n): return (datetime.date.today() + datetime.timedelta(days=n)).isoformat()
def get(path): return json.loads(urllib.request.urlopen(MOCK + path).read())
def uid(email): return sql(f"select id from profiles where email='{email}'")

ANDROID = 'Mozilla/5.0 (Linux; Android 14; Pixel 8) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Mobile Safari/537.36'
IPHONE = 'Mozilla/5.0 (iPhone; CPU iPhone OS 17_5 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.5 Mobile/15E148 Safari/604.1'

def fake_push(dev):
    """Headless Chromium can't reach Google's push service, so this device gets a stand-in subscription
    with real keys (made by the test server, which decrypts whatever is sent to it)."""
    return """(() => {
      const dev = %s, KEY = 'fake-push-sub';
      // headless Chromium always says "denied"; a real phone asks the person
      Object.defineProperty(Notification, 'permission', { get: () => 'granted' });
      Notification.requestPermission = async () => 'granted';
      const fake = o => ({
        endpoint: o.endpoint,
        options: { applicationServerKey: Uint8Array.from(o.key).buffer, userVisibleOnly: true },
        toJSON() { return { endpoint: o.endpoint, keys: { p256dh: o.p256dh, auth: o.auth } }; },
        async unsubscribe() { localStorage.removeItem(KEY); return true; },
      });
      PushManager.prototype.getSubscription = async function () { const s = localStorage.getItem(KEY); return s ? fake(JSON.parse(s)) : null; };
      PushManager.prototype.subscribe = async function (opts) {
        const o = { ...dev, key: Array.from(new Uint8Array(opts.applicationServerKey)) };
        localStorage.setItem(KEY, JSON.stringify(o));
        return fake(o);
      };
    })();""" % json.dumps(dev)

async def page_for(b, email, w=1360, h=860, ua=None, push=None, perms=()):
    ctx = await b.new_context(viewport={'width': w, 'height': h}, timezone_id='Asia/Colombo', accept_downloads=True,
                              **({'user_agent': ua} if ua else {}))
    if perms: await ctx.grant_permissions(list(perms), origin=BASE)
    if push: await ctx.add_init_script(fake_push(push))
    pg = await ctx.new_page()
    pg.on('pageerror', lambda e: errors.append(f'{email}: {e}'))
    pg.on('console', lambda m: errors.append(f'{email}: {m.text}') if m.type == 'error' and 'Failed to load resource' not in m.text else None)
    await pg.goto(BASE); await pg.fill('input[type=email]', email); await pg.fill('input[type=password]', 'password1')
    await pg.click('button:has-text("Sign in")'); await pg.wait_for_timeout(1500)
    return pg

async def to_finance(h):
    if 'Finance' not in await h.locator('.dept-switch .switcher').inner_text():
        await h.click('.dept-switch .switcher'); await h.click('.drop-item:has-text("Finance")'); await h.wait_for_timeout(1000)

async def toast(pg):
    try:
        await pg.wait_for_selector('.toast', timeout=4000)
        return ' | '.join(await pg.locator('.toast').all_inner_texts())
    except Exception:
        return ''

async def saved(dl_info, name):
    path = os.path.join(OUT, name)
    await (await dl_info.value).save_as(path)
    return path

async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()

        # =============== install as an app ===============
        h = await page_for(b, 'hasith@demo.lk', perms=('notifications', 'clipboard-read', 'clipboard-write'))
        await to_finance(h)
        check(await h.locator('link[rel=manifest]').count() == 1, 'the page links its app manifest')
        man = await h.evaluate("fetch('/manifest.webmanifest').then(r => r.json())")
        check(man['display'] == 'standalone' and man['name'] == 'Workspace' and any(i.get('purpose') == 'maskable' for i in man['icons']),
              'manifest: installable, standalone, with a maskable icon')
        sizes = await h.evaluate("Promise.all(['/icons/icon-192.png','/icons/icon-512.png','/icons/apple-touch-icon.png'].map(u => fetch(u).then(r => r.ok && r.headers.get('content-type'))))")
        check(all(s and 'png' in s for s in sizes), 'app icons are served')
        sw = await h.evaluate("navigator.serviceWorker.ready.then(r => r.active && r.active.scriptURL)")
        check(sw and sw.endswith('/sw.js'), 'the service worker is running')
        await h.click('.usermenu .icon-btn'); await h.wait_for_timeout(300)
        menu = await h.locator('.user-pop').inner_text()
        check('Install app' in menu and 'Phone alerts on this device' in menu, 'account menu offers Install app and Phone alerts')
        await h.click('.menu-list button:has-text("Install app")'); await h.wait_for_timeout(300)
        check('Install Workspace' in await h.locator('.install-help').inner_text(), 'without a browser prompt it explains how to install')
        await h.keyboard.press('Escape'); await h.mouse.click(700, 600)

        # works offline once opened (the app itself is kept on the device)
        ctx = h.context
        await ctx.set_offline(True)
        await h.reload(); await h.wait_for_timeout(1500)
        title = await h.title()
        check('Workspace' in title and await h.locator('#root').inner_html() != '', 'offline, the installed app still opens')
        await ctx.set_offline(False)
        await h.reload(); await h.wait_for_timeout(1500)
        await to_finance(h)

        ip = await page_for(b, 'amaya@demo.lk', 390, 844, ua=IPHONE)
        await ip.click('.usermenu .icon-btn'); await ip.wait_for_timeout(500)
        check('Install first' in await ip.locator('.menu-toggle').inner_text(), 'iPhone in Safari: alerts need the app installed first')
        await ip.click('.menu-list button:has-text("Install app")'); await ip.wait_for_timeout(300)
        steps = await ip.locator('.install-help').inner_text()
        check('Add to Home Screen' in steps and 'Share' in steps, 'iPhone: shows the Share → Add to Home Screen steps')
        await ip.screenshot(path=f'{OUT}/p1_iphone_install.png')

        # =============== phone alerts: admin set-up ===============
        await h.goto(BASE + '/#/admin/alerts'); await h.wait_for_timeout(1200)
        check('Set up' in await h.locator('.setup-steps').inner_text(), 'admin sees the three set-up steps')
        await h.click('button:has-text("Copy sender code")'); await h.wait_for_timeout(600)
        clip = await h.evaluate('navigator.clipboard.readText()')
        check('Deno.serve(handle)' in clip and 'aes128gcm' in clip and 'private_key' in clip, 'copies the Edge Function code to paste into Supabase')
        await h.click('button:has-text("Turn on phone alerts")'); await h.wait_for_timeout(1500)
        check(await h.locator('h2:has-text("Test it")').count() == 1, 'turning it on makes the keys and shows the test card')
        check(sql("select count(*) from workspace_push_config where enabled and length(private_key) = 43") == '1', 'the signing key is stored in the database')
        r = await h.evaluate("fetch('%s/rest/v1/workspace_push_config?select=*', {headers:{apikey:'x', Authorization:'Bearer '+JSON.parse(localStorage.getItem(Object.keys(localStorage).find(k=>k.includes('auth-token')))).access_token}}).then(r => r.status)" % MOCK)
        check(r != 200, f'…and the browser can\'t read it back ({r})')
        await h.screenshot(path=f'{OUT}/p2_admin_alerts.png')

        # admin's own phone
        hd = get('/__test/push-device')
        h2 = await page_for(b, 'hasith@demo.lk', 390, 844, ua=ANDROID, push=hd, perms=('notifications',))
        await h2.goto(BASE + '/#/admin/alerts'); await h2.wait_for_timeout(1200)
        await h2.click('button:has-text("Turn on for this device")'); await h2.wait_for_timeout(1200)
        check('This device: alerts on' in await h2.locator('.card:has-text("Test it")').inner_text(), 'admin turns alerts on for his phone')
        await h2.click('button:has-text("Send a test to my devices")')
        await h2.wait_for_selector('.push-result.good', timeout=15000)
        check('Sent to 1 device' in await h2.locator('.push-result').inner_text(), 'the test goes through the real sender code')
        inbox = get('/__test/push-inbox?endpoint=' + urllib.parse.quote(hd['endpoint']))
        check(len(inbox) == 1 and inbox[0]['title'] == 'Phone alerts are working', 'the phone decrypts the test alert')
        await h2.screenshot(path=f'{OUT}/p3_test_sent.png')

        # =============== a member's phone gets real alerts ===============
        kd = get('/__test/push-device')
        k = await page_for(b, 'kasun@demo.lk', 390, 844, ua=ANDROID, push=kd, perms=('notifications',))
        await k.click('.usermenu .icon-btn'); await k.wait_for_timeout(500)
        check('Off' in await k.locator('.menu-toggle').inner_text(), 'member: alerts start off')
        await k.click('.menu-toggle'); await k.wait_for_timeout(1500)
        check('Phone alerts are on' in await toast(k), 'member turns alerts on from the account menu')
        await k.click('.usermenu .icon-btn'); await k.wait_for_timeout(300)
        await k.click('.usermenu .icon-btn'); await k.wait_for_timeout(500)
        check('On' in await k.locator('.menu-toggle').inner_text(), 'the menu now says On')
        await k.mouse.click(200, 700)
        check(sql(f"select device from workspace_push_subscriptions where user_id='{uid('kasun@demo.lk')}'") == 'Android phone', 'his phone is listed as "Android phone"')

        await h.goto(BASE + '/#/tasks/list'); await h.wait_for_timeout(900)
        await h.fill('.qa-title', 'Send TT copy to bank')
        await h.locator('.quick-add select').select_option(label='Kasun Silva')
        await h.press('.qa-title', 'Enter'); await h.wait_for_timeout(2500)
        got = get('/__test/push-inbox?endpoint=' + urllib.parse.quote(kd['endpoint']))
        hit = [m for m in got if 'Send TT copy to bank' in m['body']]
        check(len(hit) == 1 and hit[0]['title'] == 'New task' and hit[0]['notice'], f'assigning a task buzzes his phone ({got[-1:]})')
        notice = hit[0]['notice'] if hit else 0
        tid = sql("select id from tasks where title='Send TT copy to bank'")
        check(hit and hit[0]['tag'] == f'tasks-{tid}', 'alerts about the same task replace each other on the phone')

        # tapping the alert opens the task
        await k.goto(BASE + f'/#/go?notice={notice}'); await k.wait_for_timeout(2000)
        dtxt = (await k.locator('.drawer').inner_text())[:120] if await k.locator('.drawer').count() else k.url
        check(await k.locator('.drawer .title-input').count() == 1 and await k.locator('.drawer .title-input').input_value() == 'Send TT copy to bank',
              f'tapping the alert opens that task ({dtxt!r})')
        check(sql(f"select read_at is not null from notifications where id={notice or 0}") == 't', '…and marks it read')
        await k.click('.ack-banner button:has-text("Got it")'); await k.wait_for_timeout(700)
        await k.keyboard.press('Escape')

        # the service worker shows a pushed alert
        try:
            cdp = await k.context.new_cdp_session(k)
            regs = []
            cdp.on('ServiceWorker.workerRegistrationUpdated', lambda e: regs.extend(e['registrations']))
            await cdp.send('ServiceWorker.enable'); await k.wait_for_timeout(800)
            await k.evaluate("window.__pushed = 0; window.addEventListener('workspace:pushed', () => window.__pushed++)")
            rid = [r for r in regs if r['scopeURL'].startswith(BASE)][0]['registrationId']
            await cdp.send('ServiceWorker.deliverPushMessage', {'origin': BASE, 'registrationId': rid,
                           'data': json.dumps({'title': 'Missed deadline', 'body': 'Bank rec was due', 'notice': notice, 'tag': 'x'})})
            await k.wait_for_timeout(1200)
            # (headless Chromium can't draw notifications, but the app being told proves the service worker handled it)
            check(await k.evaluate('window.__pushed') >= 1, 'the service worker receives the alert and the open app refreshes its bell')
        except Exception as e:
            check(False, f'the phone shows the alert ({e})')

        # =============== Got it ===============
        await k.goto(BASE + '/#/tasks/new'); await k.wait_for_timeout(1200)
        row = k.locator('.new-card:has-text("Clear unreconciled items list")')
        check(await row.count() == 1, 'a task he has not opened waits in his New tab')
        await k.click('.bell .icon-btn'); await k.wait_for_timeout(300)
        check("haven't opened yet" in await k.locator('.bell-pop').inner_text(), 'next morning he is reminded about unopened work')
        await k.keyboard.press('Escape'); await k.mouse.click(200, 700)
        await row.locator('.new-title').click(); await k.wait_for_selector('.drawer .title-input')
        check(await k.locator('.ack-banner').count() == 1, 'opening it asks him to press Got it')
        await k.screenshot(path=f'{OUT}/p3b_got_it.png')
        await k.click('.ack-banner button:has-text("Got it")'); await k.wait_for_timeout(800)
        check(await k.locator('.ack-banner').count() == 0, 'Got it clears the banner')
        await k.click('.subtabs button:has-text("History")'); await k.wait_for_timeout(300)
        check('opened it' in await k.locator('.history').inner_text(), 'the history shows when he opened it')
        await k.keyboard.press('Escape')
        n = await page_for(b, 'nadeesha@demo.lk')
        await n.click('.bell .icon-btn'); await n.wait_for_timeout(300)
        check('Kasun Silva hasn\'t opened "Clear unreconciled items list" yet' in await n.locator('.bell-pop').inner_text(),
              'whoever gave it heard he had not opened it')
        await n.keyboard.press('Escape')
        await n.goto(BASE + '/#/tasks/list?who=byme'); await n.wait_for_timeout(900)
        await n.click('.task-row:has-text("Clear unreconciled items list")'); await n.wait_for_selector('.drawer')
        check('opened by Kasun' in await n.locator('.drawer .meta').inner_text(), 'she can see when he opened it')
        await n.keyboard.press('Escape')

        # =============== sign-off ===============
        await k.goto(BASE + '/#/tasks/list'); await k.wait_for_timeout(900)
        await k.click('.task-row:has-text("Send TT copy to bank")'); await k.wait_for_selector('.drawer')
        await k.click('.drawer .seg-btn:has-text("Done")'); await k.wait_for_timeout(1000)
        check('Sent to Hasith for sign-off' in await toast(k), 'marking it done sends it to Hasith')
        check('waiting for' in (await k.locator('.review-banner').inner_text()).lower(), 'the owner sees it is waiting for sign-off')
        await k.keyboard.press('Escape')
        await h.goto(BASE + '/#/tasks/list'); await h.reload(); await h.wait_for_timeout(1500)
        sg = h.locator('.signoff-group')
        check(await sg.locator('.task-row:has-text("Send TT copy to bank")').count() == 1, 'it lands in "Waiting for your sign-off"')
        await sg.locator('.task-row:has-text("Send TT copy to bank")').click(); await h.wait_for_selector('.review-banner')
        await h.click('.review-banner button:has-text("Send back")'); await h.wait_for_timeout(200)
        await h.fill('.send-back textarea', 'Attach the bank acknowledgement too')
        await h.click('.send-back button:has-text("Send back")'); await h.wait_for_timeout(1500)
        check('Sent back to Kasun' in await toast(h), 'the checker sends it back with a reason')
        await h.keyboard.press('Escape')
        got = get('/__test/push-inbox?endpoint=' + urllib.parse.quote(kd['endpoint']))
        check(any(m['title'] == 'Sent back to you' and 'Attach the bank acknowledgement too' in m['body'] for m in got), 'his phone says what to fix')
        await k.reload(); await k.wait_for_timeout(1500)
        await k.click('.task-row:has-text("Send TT copy to bank")'); await k.wait_for_selector('.drawer'); await k.wait_for_timeout(600)
        check('Attach the bank acknowledgement too' in await k.locator('.sentback-banner').inner_text(), 'the task shows why it was sent back')
        await k.click('.drawer .seg-btn:has-text("Done")'); await k.wait_for_timeout(1000)
        await k.keyboard.press('Escape')
        await h.reload(); await h.wait_for_timeout(1500)
        await h.locator('.signoff-group .task-row:has-text("Send TT copy to bank")').click(); await h.wait_for_selector('.review-banner')
        await h.click('.review-banner button:has-text("Sign off")'); await h.wait_for_timeout(1000)
        check('Signed off' in await toast(h), 'then he signs it off')
        check('signed off by Hasith' in await h.locator('.drawer .meta').inner_text() and 'sent back 1×' in await h.locator('.drawer .meta').inner_text(),
              'the task records who signed it off and how often it went back')
        await h.screenshot(path=f'{OUT}/p4_signed_off.png')
        await h.keyboard.press('Escape')

        # sign-off not needed + early reminder
        await h.goto(BASE + '/#/tasks/list'); await h.wait_for_timeout(800)
        await h.fill('.qa-title', 'Print cheque register')
        await h.locator('.quick-add select').select_option(label='Kasun Silva')
        await h.press('.qa-title', 'Enter'); await h.wait_for_timeout(900)
        await h.goto(BASE + '/#/tasks/list?who=byme'); await h.wait_for_timeout(900)
        check(await h.locator('.task-row:has-text("Print cheque register") .pill.unseen').count() == 1, 'his list shows "Not opened yet"')
        await h.click('.task-row:has-text("Print cheque register")'); await h.wait_for_selector('.drawer')
        await h.locator('.drawer label:has-text("Sign-off") select').select_option('no')
        await h.locator('.drawer label:has-text("Early reminder") select').select_option('3')
        await h.click('button:has-text("Save changes")'); await h.wait_for_timeout(900)
        check(sql("select needs_check::text || remind_days from tasks where title='Print cheque register'") == 'false3', 'he switches sign-off off and sets a reminder 3 days before')
        await h.keyboard.press('Escape')
        await k.goto(BASE + '/#/tasks/new'); await k.reload(); await k.wait_for_timeout(1500)
        await k.click('.new-card:has-text("Print cheque register") .new-title'); await k.wait_for_selector('.drawer')
        check(await k.locator('.drawer label:has-text("Sign-off") select').is_disabled(), "the owner can't change the sign-off setting")
        await k.click('.ack-banner button:has-text("Got it")'); await k.wait_for_timeout(700)
        await k.keyboard.press('Escape')
        await k.goto(BASE + '/#/tasks/list'); await k.wait_for_timeout(900)
        await k.locator('.task-row:has-text("Print cheque register") .check').click(); await k.wait_for_timeout(900)
        check('Marked done' in await toast(k), 'no sign-off needed: done is done')

        # =============== Compliance calendar ===============
        await h.goto(BASE + '/#/admin/departments'); await h.wait_for_timeout(1200)
        fin = h.locator('.dept-card').filter(has=h.locator('.dept-card-head:has-text("Finance")'))
        await fin.locator('.feature-row:has-text("Compliance calendar") input').click(); await h.wait_for_timeout(1000)
        await h.goto(BASE + '/#/tasks/x-compliance'); await h.wait_for_timeout(1200)
        await to_finance(h)
        check(await h.locator('.tab.active').inner_text() == 'Compliance', 'Finance gets a Compliance tab once switched on')
        await h.click('button:has-text("+ Add obligation")')
        await h.fill('.co-add input[list=co-names]', 'VAT return')
        await h.fill('.co-add input[placeholder="e.g. IRD"]', 'IRD')
        await h.locator('.co-add label:has-text("Owner") select').select_option(label='Kasun Silva')
        await h.fill('.co-add input[type=date]', day(12))
        await h.locator('.co-add label:has-text("Early reminder") select').select_option('5')
        await h.click('.co-add button:has-text("Add obligation")'); await h.wait_for_timeout(1500)
        check('Kasun has been given the task' in await toast(h), 'adding an obligation gives the owner its task')
        row = h.locator('.co-table tr:has-text("VAT return")')
        txt = await row.inner_text()
        check('12d left' in txt and 'Kasun' in txt and 'IRD' in txt and 'Monthly' in txt, f'the row shows deadline, days left, owner ({txt!r})')
        check(sql("select priority || remind_days || repeat from tasks where title='VAT return'") == 'high5monthly', 'it is a high-priority monthly task with a 5-day reminder')
        await h.screenshot(path=f'{OUT}/p5_compliance.png')
        await h.click('.seg-btn:has-text("Next 12 months")'); await h.wait_for_timeout(500)
        months = await h.locator('.co-month').count()
        vats = await h.locator('.co-month li:has-text("VAT return")').count()
        check(months == 12 and vats in (11, 12), f'the year view shows every month with the schedule ({months} months, {vats} VAT dates)')
        await h.screenshot(path=f'{OUT}/p6_compliance_year.png')
        await h.click('.seg-btn:has-text("List")'); await h.wait_for_timeout(300)
        dl = h.expect_download()
        async with dl as info:
            await h.click('button:has-text("Export to Excel")')
        wb = openpyxl.load_workbook(await saved(info, 'compliance.xlsx'))
        ws = wb['Compliance']
        check(ws.cell(1, 1).value == 'Obligation' and ws.cell(2, 1).value == 'VAT return' and ws.cell(2, 4).value == 'Kasun Silva', 'exports to Excel')
        await k.goto(BASE + '/#/tasks/x-compliance'); await k.reload(); await k.wait_for_timeout(1500)
        check(await k.locator('.co-table tr:has-text("VAT return")').count() == 1 and await k.locator('button:has-text("+ Add obligation")').count() == 0,
              'members see the calendar but cannot add to it')
        pr = await page_for(b, 'priya@demo.lk')
        check(await pr.locator('.tab:has-text("Compliance")').count() == 0, 'HR has no Compliance tab')

        # =============== Trends ===============
        await h.goto(BASE + '/#/tasks/week'); await h.wait_for_timeout(1000)
        check(await h.locator('.tab.active').inner_text() == 'Reports', 'the weekly summary now lives under Reports')
        await h.click('.seg-btn:has-text("Trends")'); await h.wait_for_timeout(1500)
        check(await h.locator('h1').inner_text() == 'Trends' and await h.locator('.trend-table tbody tr').count() >= 3, 'Trends shows a row per person')
        check('Everyone' in await h.locator('.trend-table').inner_text(), '…and a department total')
        await h.click('.trend-metrics .seg-btn:has-text("Time to open")'); await h.wait_for_timeout(500)
        check('pressing "Got it"' in await h.locator('main').inner_text(), 'each measure explains itself')
        await h.screenshot(path=f'{OUT}/p7_trends.png')
        dl = h.expect_download()
        async with dl as info:
            await h.click('button:has-text("Export to Excel")')
        ws = openpyxl.load_workbook(await saved(info, 'trends.xlsx'))['Trends']
        check(ws.cell(1, 1).value == 'Person' and ws.cell(1, 7).value == 'On time %' and ws.max_row > 2, 'trends export to Excel')
        await k.goto(BASE + '/#/tasks/trends'); await k.wait_for_timeout(1000)
        check(await k.locator('.tab:has-text("Reports")').count() == 0 and await k.locator('h1:has-text("Trends")').count() == 0, 'members do not get Reports')

        # =============== Search ===============
        await k.goto(BASE + '/#/tasks/list'); await k.wait_for_timeout(1000)
        await k.click('.search-btn'); await k.wait_for_timeout(300)
        await k.fill('.search-input input', 'petty'); await k.wait_for_timeout(1200)
        check(await k.locator('.search-hit:has-text("Petty cash count")').count() == 1, 'search finds a task by title')
        await k.keyboard.press('Enter'); await k.wait_for_timeout(1200)
        check(await k.locator('.drawer .title-input').input_value() == 'Petty cash count & top-up', 'Enter opens it')
        await k.keyboard.press('Escape')
        await h.goto(BASE + '/#/tasks/list'); await h.wait_for_timeout(900)
        await h.keyboard.press('Control+k'); await h.wait_for_timeout(300)
        await h.fill('.search-input input', 'wrong period'); await h.wait_for_timeout(1200)
        res = await h.locator('.search-results').inner_text()
        check('Review VAT return schedules' in res and 'Comment:' in res, 'Ctrl+K finds words inside comments')
        await h.keyboard.press('Escape')
        # a note page with a word in its text
        hid = uid('hasith@demo.lk')
        sql(f"""begin; set local role authenticated; select set_config('request.jwt.claims', '{{"sub":"{hid}"}}', true);
            insert into notes_pages (title, content, share_scope) values ('Audit queries',
            '{{"type":"doc","content":[{{"type":"paragraph","content":[{{"type":"text","text":"Ask auditors about the depreciation rerun"}}]}}]}}', 'department'); commit;""")
        await h.keyboard.press('Control+k'); await h.wait_for_timeout(300)
        await h.fill('.search-input input', 'depreciation'); await h.wait_for_timeout(1200)
        check(await h.locator('.search-results section:has(h4:has-text("Notes")) .search-hit:has-text("Audit queries")').count() == 1, 'search finds words inside note pages')
        await h.screenshot(path=f'{OUT}/p8_search.png')
        await h.keyboard.press('Escape')
        await pr.keyboard.press('Control+k'); await pr.wait_for_timeout(300)
        await pr.fill('.search-input input', 'petty'); await pr.wait_for_timeout(1200)
        check('Nothing found' in await pr.locator('.search-results').inner_text(), "HR's manager can't find Finance's work")

        # =============== phone alerts admin view + sign-out ===============
        await h.goto(BASE + '/#/admin/alerts'); await h.wait_for_timeout(1500)
        page = await h.locator('main').inner_text()
        check('Kasun Silva' in page and 'Android phone' in page and 'Recent sends' in page, 'admin sees who has alerts on and recent sends')
        await k.click('.usermenu .icon-btn'); await k.click('.menu-list button:has-text("Sign out")'); await k.wait_for_timeout(1500)
        check(sql(f"select count(*) from workspace_push_subscriptions where user_id='{uid('kasun@demo.lk')}'") == '0', 'signing out stops alerts on that phone')

        # phones: no sideways scrolling on the new pages
        m = await page_for(b, 'hasith@demo.lk', 390, 844)
        await to_finance(m)
        for path in ['/#/tasks/x-compliance', '/#/tasks/trends', '/#/admin/alerts']:
            await m.goto(BASE + path); await m.wait_for_timeout(1000)
            sw = await m.evaluate('document.documentElement.scrollWidth')
            check(sw <= 392, f'{path} fits a phone screen ({sw}px)')
        await m.goto(BASE + '/#/tasks/list'); await m.wait_for_timeout(800)
        await m.click('.search-btn'); await m.wait_for_timeout(300)
        await m.screenshot(path=f'{OUT}/p9_mobile_search.png')
        await b.close()
    print('ERRORS:', errors or 'none')
    print(sum(results), 'passed,', len(results) - sum(results), 'failed')

if __name__ == '__main__':
    asyncio.run(main())
